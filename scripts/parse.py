# -*- coding: utf-8 -*-
"""
여러 양식의 카드·은행·가계부 파일을 하나의 표준 거래표로 합친다.

  python -X utf8 scripts/parse.py                      (기본: data/sample)
  python -X utf8 scripts/parse.py --input data/private
  python -X utf8 scripts/parse.py --input "다른폴더" --output "out/내결과.csv"

사람 이름은 폴더 이름에서 가져온다. 파일 내용으로 사람을 추측하지 않는다.
"""
import argparse
import csv
import datetime as dt
import sys
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path

# ============================================================================
# CONFIG — 고칠 곳은 여기 한 군데뿐입니다
# ============================================================================

# 파일마다 컬럼 이름이 다릅니다. 내 파일의 컬럼 이름을 여기에 추가하세요.
컬럼별칭 = {
    "날짜":     ["날짜", "거래일자", "거래일시", "이용일자", "승인일자", "승인일시",
                "사용일", "사용일자", "거래날짜", "일자"],
    "가맹점":   ["가맹점", "가맹점명", "이용하신곳", "이용가맹점", "내용", "적요",
                "사용처", "상호", "거래내용", "메모", "설명"],
    "금액":     ["금액", "이용금액", "거래금액", "사용금액", "결제금액", "승인금액"],
    "출금":     ["출금액", "출금금액", "출금", "지급액", "차변"],
    "입금":     ["입금액", "입금금액", "입금", "예입액", "대변"],
    "구분":     ["구분", "수입지출", "타입", "거래구분", "입출금구분"],
    "원본분류": ["분류", "카테고리", "업종", "업종구분", "대분류", "소분류", "유형"],
    "결제수단": ["결제수단", "결제방법", "카드명", "지급수단", "할부", "할부개월"],
}

# 가맹점 이름에 이 글자가 들어 있으면 그 카테고리로 봅니다 (위에서부터 먼저 검사).
카테고리규칙 = [
    ("구독",   ["넷플릭스", "netflix", "왓챠", "티빙", "웨이브", "디즈니", "유튜브 프리미엄",
              "멜론", "스포티파이", "쿠팡플레이", "지니뮤직", "챗gpt", "구독"]),
    ("통신비", ["skt", "kt ", "lg유플러스", "lgu+", "통신요금", "알뜰폰", "통신비"]),
    ("주거",   ["관리비", "아파트", "월세", "전기요금", "도시가스", "수도요금"]),
    ("보험",   ["생명", "화재", "해상", "손해보험", "보험료", "실손"]),
    ("교육",   ["학원", "교습", "학습지", "등록금", "수업료", "교육비"]),
    ("운동",   ["헬스", "짐", "요가", "필라테스", "수영장", "골프"]),
    ("의료",   ["병원", "의원", "약국", "치과", "한의원", "클리닉", "메디컬", "진료"]),
    ("교통",   ["택시", "카카오t", "주유", "칼텍스", "오일뱅크", "sk에너지", "지하철",
              "버스", "코레일", "하이패스", "주차"]),
    ("장보기", ["이마트", "홈플러스", "롯데마트", "트레이더스", "코스트코", "컬리",
              "농협하나로", "마트", "정육", "청과"]),
    ("식비",   ["스타벅스", "커피", "카페", "베이커리", "파리바게", "뚜레쥬르", "투썸",
              "이디야", "배달의민족", "요기요", "쿠팡이츠", "김밥", "도시락", "버거",
              "피자", "치킨", "식당", "음식", "고기", "중화", "분식", "편의점",
              "gs25", "cu ", "세븐일레븐", "이마트24"]),
    ("쇼핑",   ["무신사", "올리브영", "백화점", "아울렛", "지그재그", "에이블리",
              "유니클로", "자라", "나이키", "아디다스"]),
    ("문화",   ["cgv", "메가박스", "롯데시네마", "교보문고", "예스24", "알라딘",
              "영화", "서점", "공연", "전시"]),
    ("생활",   ["다이소", "쿠팡", "11번가", "지마켓", "옥션", "네이버페이", "생활용품",
              "드럭스토어", "문구"]),
]

# 원본 파일에 분류 컬럼이 있을 때 표준 카테고리로 바꿉니다 (위 규칙이 먼저입니다).
원본분류매핑 = {
    "카페": "식비", "음식점": "식비", "배달": "식비", "제과": "식비", "외식": "식비",
    "대형마트": "장보기", "식료품": "장보기",
    "온라인쇼핑": "생활", "생활용품": "생활",
    "의류": "쇼핑", "화장품": "쇼핑", "백화점": "쇼핑",
    "병원": "의료", "약국": "의료",
    "스포츠": "운동", "주유": "교통", "서점": "문화", "선물": "쇼핑",
}

# 가맹점·적요에 이 글자가 있으면 '이체'로 보고 소비 합계에서 뺍니다.
# (가족끼리 주고받은 돈, 카드대금 결제 — 두 번 세지 않기 위함)
이체키워드 = ["카드대금", "카드값", "용돈", "계좌이체", "내계좌", "자동이체 반환", "송금"]

# 수입으로 보는 글자
수입키워드 = ["급여", "월급", "상여", "보너스", "환급", "이자", "배당"]

# 파일을 읽을 때 시도하는 인코딩 순서 (앞에서부터 하나씩 해봅니다)
인코딩후보 = ["utf-8-sig", "utf-8", "cp949"]

# 날짜 형식 후보
날짜형식 = ["%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d", "%Y%m%d",
          "%Y-%m-%d %H:%M:%S", "%Y.%m.%d %H:%M", "%Y-%m-%d %H:%M"]

# ============================================================================
# 여기서부터는 건드리지 않아도 됩니다
# ============================================================================

표준컬럼 = ["날짜", "사람", "원본파일", "구분", "가맹점", "금액", "카테고리", "결제수단", "중복의심"]


class 읽기실패(Exception):
    """사용자에게 무엇을 고쳐야 하는지 알려주는 오류."""


# ----------------------------------------------------------------- 인코딩
def 인코딩찾기(path: Path):
    """strict 디코딩만 시도한다. 글자를 버리거나 바꿔서 성공 처리하지 않는다."""
    raw = path.read_bytes()
    시도 = []
    for enc in 인코딩후보:
        try:
            text = raw.decode(enc)
        except UnicodeDecodeError as e:
            시도.append(f"{enc} 실패({e.reason})")
            continue
        return text, enc, 시도
    raise 읽기실패(
        f"{path.name} — 글자 인코딩을 알 수 없습니다.\n"
        f"    시도한 것: {' / '.join(시도)}\n"
        f"    고칠 곳: 엑셀에서 이 파일을 열고 'CSV UTF-8'로 다시 저장해 주세요.\n"
        f"    또는 scripts/parse.py 의 CONFIG > 인코딩후보 에 인코딩을 추가하세요."
    )


# ----------------------------------------------------------------- 표 읽기
def csv_행읽기(path: Path):
    text, enc, _ = 인코딩찾기(path)
    rows = list(csv.reader(text.splitlines()))
    return rows, enc


def xlsx_행읽기(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise 읽기실패(
            f"{path.name} — 엑셀 파일을 읽으려면 openpyxl 이 필요합니다.\n"
            f"    고칠 곳: 명령창에 python -m pip install openpyxl 을 입력하세요."
        )
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise 읽기실패(
            f"{path.name} — 엑셀 파일을 열지 못했습니다. ({type(e).__name__})\n"
            f"    고칠 곳: 엑셀에서 열어 .xlsx 형식으로 다시 저장해 주세요.\n"
            f"    (옛날 .xls 형식이나 암호가 걸린 파일은 읽을 수 없습니다.)"
        )
    ws = wb[wb.sheetnames[0]]
    rows = [["" if c is None else c for c in r] for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows, "xlsx"


# ----------------------------------------------------------------- 헤더 찾기
def 정규화(s) -> str:
    if s is None:
        return ""
    return unicodedata.normalize("NFC", str(s)).strip()


def 헤더행찾기(rows):
    """알려진 컬럼 이름이 2개 이상 나오는 첫 행을 헤더로 본다."""
    알려진 = {정규화(a) for names in 컬럼별칭.values() for a in names}
    for i, row in enumerate(rows[:30]):
        cells = {정규화(c) for c in row if 정규화(c)}
        if len(cells & 알려진) >= 2:
            return i
    return None


def 컬럼매핑(header):
    """헤더 → {표준이름: 열번호}"""
    매핑 = {}
    정규화된 = [정규화(h) for h in header]
    for 표준, 후보들 in 컬럼별칭.items():
        for 후보 in 후보들:
            후보n = 정규화(후보)
            for idx, h in enumerate(정규화된):
                if h == 후보n and 표준 not in 매핑:
                    매핑[표준] = idx
    return 매핑


# ----------------------------------------------------------------- 값 정규화
def 날짜읽기(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = 정규화(v)
    if not s:
        return None
    for fmt in 날짜형식:
        try:
            return dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def 금액읽기(v):
    """'8,500원' → 8500, '-45,000' → -45000. 읽을 수 없으면 None."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(round(float(v)))
    s = 정규화(v)
    if not s:
        return None
    음수 = s.startswith("-") or (s.startswith("(") and s.endswith(")"))
    for 버릴것 in [",", "원", "₩", "KRW", "krw", " ", "(", ")", "-", "+"]:
        s = s.replace(버릴것, "")
    if not s:
        return None
    try:
        값 = int(Decimal(s))
    except (InvalidOperation, ValueError):
        return None
    return -값 if 음수 else 값


def 카테고리정하기(가맹점: str, 원본분류: str) -> str:
    이름 = 가맹점.lower()
    for 카테고리, 키워드들 in 카테고리규칙:
        for kw in 키워드들:
            if kw.lower() in 이름:
                return 카테고리
    분류 = 정규화(원본분류)
    if 분류 in 원본분류매핑:
        return 원본분류매핑[분류]
    if 분류 in {c for c, _ in 카테고리규칙}:
        return 분류
    return "기타"


def 구분정하기(가맹점, 구분값, 금액, 수입가능):
    """수입가능=False 인 파일(카드 이용내역 등)에서는 금액이 양수여도 수입으로 보지 않는다.
    카드 내역의 양수 금액은 환불이므로 '지출'로 두어야 그달 지출에서 차감된다."""
    이름 = 가맹점.lower()
    for kw in 이체키워드:
        if kw.lower() in 이름:
            return "이체"
    구분n = 정규화(구분값)
    if 구분n in ("수입", "입금"):
        return "수입"
    if 구분n in ("지출", "출금"):
        return "지출"
    if not 수입가능:
        return "지출"
    for kw in 수입키워드:
        if kw.lower() in 이름:
            return "수입"
    return "수입" if 금액 > 0 else "지출"


# ----------------------------------------------------------------- 파일 1개 처리
def 파일읽기(path: Path, 사람: str, 루트: Path):
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        rows, enc = xlsx_행읽기(path)
    elif path.suffix.lower() == ".csv":
        rows, enc = csv_행읽기(path)
    else:
        return [], enc_none(path)

    if not rows:
        raise 읽기실패(f"{path.name} — 파일이 비어 있습니다.")

    h = 헤더행찾기(rows)
    if h is None:
        raise 읽기실패(
            f"{path.name} — 표의 제목 줄(헤더)을 찾지 못했습니다.\n"
            f"    파일 첫 줄: {' | '.join(정규화(c) for c in rows[0][:8])}\n"
            f"    고칠 곳: scripts/parse.py 의 CONFIG > 컬럼별칭 에 이 파일의 컬럼 이름을 추가하세요."
        )

    header = rows[h]
    매핑 = 컬럼매핑(header)
    실제컬럼 = " | ".join(정규화(c) for c in header if 정규화(c))

    if "날짜" not in 매핑:
        raise 읽기실패(
            f"{path.name} — 날짜 컬럼을 찾지 못했습니다.\n"
            f"    이 파일의 실제 컬럼: {실제컬럼}\n"
            f"    고칠 곳: scripts/parse.py 의 CONFIG > 컬럼별칭 > \"날짜\" 에 위 이름 중 하나를 추가하세요."
        )
    if "가맹점" not in 매핑:
        raise 읽기실패(
            f"{path.name} — 가맹점(내용) 컬럼을 찾지 못했습니다.\n"
            f"    이 파일의 실제 컬럼: {실제컬럼}\n"
            f"    고칠 곳: scripts/parse.py 의 CONFIG > 컬럼별칭 > \"가맹점\" 에 위 이름 중 하나를 추가하세요."
        )
    if "금액" not in 매핑 and not ("출금" in 매핑 or "입금" in 매핑):
        raise 읽기실패(
            f"{path.name} — 금액 컬럼을 찾지 못했습니다.\n"
            f"    이 파일의 실제 컬럼: {실제컬럼}\n"
            f"    고칠 곳: scripts/parse.py 의 CONFIG > 컬럼별칭 > \"금액\" 에 위 이름 중 하나를 추가하세요."
        )

    상대경로 = path.relative_to(루트).as_posix()
    결과, 건너뜀 = [], 0
    # 출금/입금이 나뉘어 있거나 구분 컬럼이 있는 파일만 수입이 나올 수 있다.
    수입가능 = ("출금" in 매핑) or ("입금" in 매핑) or ("구분" in 매핑)

    for n, row in enumerate(rows[h + 1:], start=h + 2):
        def 셀(이름):
            i = 매핑.get(이름)
            return row[i] if i is not None and i < len(row) else ""

        날짜 = 날짜읽기(셀("날짜"))
        가맹점 = 정규화(셀("가맹점"))
        if not 날짜 or not 가맹점:
            건너뜀 += 1
            continue

        if "출금" in 매핑 or "입금" in 매핑:
            출 = 금액읽기(셀("출금")) or 0
            입 = 금액읽기(셀("입금")) or 0
            금액 = 입 - 출
        else:
            원금액 = 금액읽기(셀("금액"))
            if 원금액 is None:
                건너뜀 += 1
                continue
            구분값 = 정규화(셀("구분"))
            금액 = 원금액 if 구분값 in ("수입", "입금") else -원금액

        if 금액 == 0:
            건너뜀 += 1
            continue

        구분 = 구분정하기(가맹점, 셀("구분"), 금액, 수입가능)
        결과.append({
            "날짜": 날짜,
            "사람": 사람,
            "원본파일": f"{상대경로}:{n}행",
            "구분": 구분,
            "가맹점": 가맹점,
            "금액": 금액,
            "카테고리": "이체" if 구분 == "이체" else 카테고리정하기(가맹점, 셀("원본분류")),
            "결제수단": 정규화(셀("결제수단")) or "",
            "중복의심": "",
        })

    return 결과, (enc, len(결과), 건너뜀)


def enc_none(path):
    return ("건너뜀", 0, 0)


# ----------------------------------------------------------------- 중복 의심
def 중복표시(거래들):
    """같은 사람·같은 금액·날짜 차이 1일 이내 → 중복 의심. 절대 자동 삭제하지 않는다."""
    본 = {}
    의심 = 0
    for t in 거래들:
        키 = (t["사람"], t["금액"])
        d = dt.date.fromisoformat(t["날짜"])
        for 기존 in 본.get(키, []):
            if abs((d - dt.date.fromisoformat(기존["날짜"])).days) <= 1 \
                    and 기존["원본파일"].split(":")[0] != t["원본파일"].split(":")[0]:
                t["중복의심"] = "Y"
                기존["중복의심"] = "Y"
                의심 += 1
                break
        본.setdefault(키, []).append(t)
    return 의심


# ----------------------------------------------------------------- 메인
def main():
    sys.stdout.reconfigure(encoding="utf-8")

    ap = argparse.ArgumentParser(description="여러 양식의 금융 파일을 한 표로 합칩니다.")
    ap.add_argument("--input", default="data/sample", help="입력 폴더 (기본: data/sample)")
    ap.add_argument("--output", default="out/거래통합.csv", help="출력 CSV (기본: out/거래통합.csv)")
    args = ap.parse_args()

    루트 = Path(args.input).resolve()
    if not 루트.is_dir():
        print(f"[!] 입력 폴더가 없습니다: {루트}")
        return 3

    print(f"입력 폴더: {루트}")
    print()

    거래들, 읽은파일, 실패 = [], 0, []

    for 사람폴더 in sorted(p for p in 루트.iterdir() if p.is_dir()):
        사람 = unicodedata.normalize("NFC", 사람폴더.name)
        파일들 = sorted(p for p in 사람폴더.rglob("*")
                     if p.is_file() and p.suffix.lower() in (".csv", ".xlsx", ".xlsm"))
        for f in 파일들:
            try:
                rows, (enc, n, skip) = 파일읽기(f, 사람, 루트)
            except 읽기실패 as e:
                실패.append(str(e))
                continue
            거래들.extend(rows)
            읽은파일 += 1
            꼬리 = f" (건너뛴 줄 {skip})" if skip else ""
            print(f"  [{사람}] {f.name}  {enc}  {n}건{꼬리}")

    if 실패:
        print()
        for msg in 실패:
            print(f"[!] {msg}")

    if not 거래들:
        print("\n읽어 들인 거래가 없습니다.")
        return 3

    거래들.sort(key=lambda t: (t["날짜"], t["사람"]))
    의심 = 중복표시(거래들)

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    with open(out, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=표준컬럼)
        w.writeheader()
        w.writerows(거래들)

    # ---- 요약 ----
    달들 = sorted({t["날짜"][:7] for t in 거래들})
    사람별 = {}
    for t in 거래들:
        사람별.setdefault(t["사람"], []).append(t)

    def 지출합(ts):
        return sum(-t["금액"] for t in ts if t["구분"] == "지출")

    print()
    print(f"읽은 파일 {읽은파일}개 / 거래 {len(거래들)}건 / 기간 {달들[0]} ~ {달들[-1]}")
    print("사람별: " + " · ".join(
        f"{이름} {len(ts)}건 {지출합(ts):,}원" for 이름, ts in sorted(사람별.items())))

    수입 = sum(t["금액"] for t in 거래들 if t["구분"] == "수입")
    지출 = 지출합(거래들)
    이체 = sum(1 for t in 거래들 if t["구분"] == "이체")
    print(f"가구 총지출 {지출:,}원 / 총수입 {수입:,}원 / 이체 {이체}건 (합계에서 제외)")
    if 의심:
        print(f"중복 의심 {의심}쌍 — 자동으로 지우지 않았습니다. 대시보드에서 확인하세요.")
    if 실패:
        print(f"읽지 못한 파일 {len(실패)}개 — 위 [!] 안내를 보세요.")

    print(f"\n저장: {out}")
    return 0 if not 실패 else 2


if __name__ == "__main__":
    sys.exit(main())
